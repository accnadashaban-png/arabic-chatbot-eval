"""End to end on the real golden set: runs, gates, comparison, reports, review and the judge."""
from pathlib import Path

import httpx
import pytest
from openpyxl import load_workbook

from chatbot_eval.__main__ import main
from chatbot_eval.bots import MockBot
from chatbot_eval.builder import build_cases
from chatbot_eval.compare import ComparisonError, compare
from chatbot_eval.gates import Gates, decide
from chatbot_eval.judge import JudgeError, OpenAICompatibleModel, grade, parse_grade
from chatbot_eval.metrics import summarize
from chatbot_eval.report import write_excel, write_markdown
from chatbot_eval.review import agreement, export_review, import_review
from chatbot_eval.runner import Run, run

CASES = build_cases()


@pytest.fixture(scope="module")
def rounds() -> tuple[Run, Run]:
    return run(CASES, MockBot(1), label="round 1"), run(CASES, MockBot(2), label="round 2")


class Broken:
    name = "broken"

    def ask(self, question: str) -> str:
        if "بكام" in question:
            raise TimeoutError("no answer in 60 s")
        return "ok"


def test_a_failing_call_is_recorded_not_fatal() -> None:
    r = run(CASES[:60], Broken(), workers=3)
    errors = [x for x in r.results if x.error]
    assert errors and all(x.failed == ["error"] for x in errors) and len(r.results) == 60
    assert [x.case_id for x in r.results] == [c.id for c in CASES[:60]]  # order kept despite threads


def test_run_round_trips_through_jsonl(tmp_path: Path, rounds) -> None:
    first, _ = rounds
    loaded = Run.load(first.save(tmp_path / "r1.jsonl"))
    assert loaded.results == first.results and loaded.label == "round 1" and loaded.dataset == first.dataset


def test_round_one_is_blocked_by_leaks_round_two_by_slices(rounds) -> None:
    first, second = rounds
    s1, s2 = summarize(first, CASES), summarize(second, CASES)
    d1, d2 = decide(s1, Gates()), decide(s2, Gates())
    assert not d1.go and len(s1.critical_failures) == 8 and s1.reasons["forbidden"] == 8
    assert not d2.go and not s2.critical_failures
    assert {g.name for g in d2.blockers} == {"Overall pass rate", "Dialect: maghrebi", "Intent: pricing"}
    assert s1.slices["dialect"][0].value == "franco" and s1.slices["dialect"][0].rate == 0.0  # worst slice first


def test_gates_pass_a_good_bot_and_skip_tiny_slices(rounds) -> None:
    _, second = rounds
    lenient = Gates(min_pass_rate=0.85, min_slice_pass_rate=0.05, min_slice_size=70)
    d = decide(summarize(second, CASES), lenient)
    assert d.go and any("english" in t for t in d.too_small) and not any(g.name == "Dialect: english" for g in d.gates)


def test_comparison_counts_and_significance(rounds) -> None:
    first, second = rounds
    c = compare(first, second, CASES)
    a, b = first.by_id(), second.by_id()
    assert len(c.fixed) == sum(not a[i].passed and b[i].passed for i in a)
    assert len(c.regressed) == sum(a[i].passed and not b[i].passed for i in a)
    assert c.significant and c.deltas[0].value == "pricing" and c.deltas[0].delta < -0.5
    assert len(c.fixed) + len(c.regressed) + len(c.still_failing) + c.still_passing == len(CASES)
    other = Run("x", "b", "different", "now", second.results)
    with pytest.raises(ComparisonError, match="different golden sets"):
        compare(first, other, CASES)


def test_reports_contain_the_decision_and_evidence(tmp_path: Path, rounds) -> None:
    first, second = rounds
    s = summarize(second, CASES)
    d = decide(s, Gates())
    c = compare(first, second, CASES)
    wb = load_workbook(write_excel(tmp_path / "r.xlsx", second, CASES, s, d, c))
    assert wb.sheetnames == ["Decision", "Summary", "By dialect", "By intent", "By test type", "By difficulty", "Coverage",
                             "Failures", "Comparison", "All results"]
    assert wb["Decision"]["A1"].value == "Launch decision: NO-GO"
    assert wb["Failures"].max_row == 1 + (s.total - s.passed) and wb["All results"].max_row == 1 + s.total
    md = write_markdown(tmp_path / "r.md", s, d, c, CASES, second).read_text(encoding="utf-8")
    assert "**Decision: NO-GO**" in md and "| Intent: pricing |" in md and "Fixed **" in md


def test_review_sheet_round_trip_and_agreement(tmp_path: Path, rounds) -> None:
    _, second = rounds
    path = export_review(second, CASES, tmp_path / "review.xlsx")
    wb = load_workbook(path)
    ws = wb["Review"]
    ids = [row[0].value for row in ws.iter_rows(min_row=2)]
    auto = {r.case_id: "pass" if r.passed else "fail" for r in second.results}
    for n, row in enumerate(ws.iter_rows(min_row=2)):  # a reviewer agrees except on two rows
        row[7].value = auto[row[0].value] if n not in (0, 1) else ("fail" if auto[row[0].value] == "pass" else "pass")
    wb.save(path)
    human = import_review(path)
    ag = agreement({i: auto[i] for i in ids}, human)
    assert ag.reviewed == len(ids) == 60 and ag.accuracy == pytest.approx(58 / 60) and 0.8 < ag.kappa < 1
    assert len(ag.auto_too_lenient) + len(ag.auto_too_strict) == 2


def test_judge_parses_and_validates_grades() -> None:
    g = parse_grade("C1", 'Sure:\n```json\n{"correct": 5, "language": 4, "tone": 5, "safe": true, "verdict": "PASS", "reason": "ok"}\n```')
    assert g.verdict == "pass" and g.correct == 5
    for bad in ("no json here", '{"correct": 9, "language": 4, "tone": 5, "safe": true, "verdict": "pass"}', '{"correct": 4}'):
        with pytest.raises(JudgeError):
            parse_grade("C1", bad)


def test_judge_calls_an_openai_compatible_model() -> None:
    seen = []

    def handler(request: httpx.Request) -> httpx.Response:
        seen.append(request.read().decode())
        return httpx.Response(200, json={"choices": [{"message": {"content":
                              '{"correct": 2, "language": 5, "tone": 4, "safe": true, "verdict": "fail", "reason": "no handoff"}'}}]})
    model = OpenAICompatibleModel("https://judge.test/v1/chat/completions", "k", "m", transport=httpx.MockTransport(handler))
    refund = next(c for c in CASES if c.intent == "refund")
    assert grade(model, refund, "Refunds are possible within 14 days.").verdict == "fail"
    assert "handoff" in seen[0] and refund.question in seen[0]


def test_cli_exit_codes(tmp_path: Path) -> None:
    out = tmp_path / "r2.jsonl"
    assert main(["run", "--bot", "mock:v2", "--label", "cli", "--out", str(out)]) == 0
    assert main(["report", str(out), "--strict"]) == 1  # NO-GO fails a CI job
    assert main(["report", str(out), "--md", str(tmp_path / "r.md")]) == 0
    assert main(["validate"]) == 0 and main(["validate", "--strict"]) == 1  # the planted contamination
    assert main(["report", str(tmp_path / "missing.jsonl")]) == 2
