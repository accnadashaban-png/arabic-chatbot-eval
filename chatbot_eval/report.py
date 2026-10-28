"""Excel and Markdown reports: the decision first, then the evidence."""
from __future__ import annotations

from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.worksheet import Worksheet

from chatbot_eval.compare import Comparison
from chatbot_eval.dataset import DIALECTS, Case, coverage
from chatbot_eval.gates import Decision
from chatbot_eval.metrics import Summary
from chatbot_eval.runner import Run

HEADER = PatternFill("solid", fgColor="1F5A4F")
GOOD, WARN, BAD = (PatternFill("solid", fgColor=c) for c in ("E4EEE9", "F6EBDC", "FBEAEA"))
SHEET_TAGS = {"dialect": "By dialect", "intent": "By intent", "test_type": "By test type", "difficulty": "By difficulty"}


def _table(ws: Worksheet, header: list[str], rows: list[list[Any]], widths: list[int] | None = None) -> None:
    ws.append(header)
    for cell in ws[ws.max_row]:
        cell.font, cell.fill = Font(bold=True, color="FFFFFF"), HEADER
    for row in rows:
        ws.append(row)
    for i, w in enumerate(widths or []):
        ws.column_dimensions[chr(65 + i)].width = w
    ws.freeze_panes = ws.cell(row=ws.max_row - len(rows) + 1, column=1).coordinate


def _rate_fill(rate: float, threshold: float) -> PatternFill:
    return GOOD if rate >= threshold else WARN if rate >= threshold - 0.1 else BAD


def write_excel(path: str | Path, run: Run, cases: list[Case], s: Summary, d: Decision,
                comparison: Comparison | None = None, slice_threshold: float = 0.8) -> Path:
    wb = Workbook()
    ws = wb.active
    assert ws is not None
    ws.title = "Decision"
    ws.append([f"Launch decision: {d.label}"])
    ws["A1"].font = Font(bold=True, size=16, color="1F5A4F" if d.go else "A12C2C")
    ws.append([f"{s.label or s.run_id} · bot {s.bot} · golden set {s.dataset} · {s.total} cases"])
    ws.append([])
    _table(ws, ["Gate", "Value", "Threshold", "Result"],
           [[g.name, g.value, g.threshold, "pass" if g.ok else "FAIL"] for g in d.gates], [52, 14, 16, 10])
    for row in ws.iter_rows(min_row=5):
        row[3].fill = GOOD if row[3].value == "pass" else BAD
    if d.too_small:
        ws.append([])
        ws.append(["Not judged (too few cases to be fair): " + "; ".join(d.too_small)])

    ws = wb.create_sheet("Summary")
    lo, hi = s.ci
    _table(ws, ["Metric", "Value"], [
        ["Test cases", s.total], ["Passed", s.passed], ["Pass rate", f"{s.pass_rate:.1%}"],
        ["Pass rate, 95% interval", f"{lo:.1%} to {hi:.1%}"], ["Answer language matches", f"{s.language_rate:.1%}"],
        ["Critical failures", len(s.critical_failures)],
        *[[f"Failed: {k}", v] for k, v in s.reasons.items()],
        ["Median response", f"{s.p50_ms:.0f} ms" if s.p50_ms is not None else "–"],
        ["p95 response", f"{s.p95_ms:.0f} ms" if s.p95_ms is not None else "–"],
    ], [34, 22])

    for tag, title in SHEET_TAGS.items():
        ws = wb.create_sheet(title)
        rows = [[sl.value, sl.total, sl.passed, round(sl.rate, 3), round(sl.ci[0], 3), round(sl.ci[1], 3)] for sl in s.slices[tag]]
        _table(ws, [tag, "cases", "passed", "pass rate", "95% low", "95% high"], rows, [18, 8, 8, 10, 10, 10])
        for r in ws.iter_rows(min_row=2):
            r[3].fill = _rate_fill(float(r[3].value or 0), slice_threshold)
            r[3].number_format = r[4].number_format = r[5].number_format = "0.0%"

    ws = wb.create_sheet("Coverage")
    matrix = coverage(cases)
    _table(ws, ["intent", *DIALECTS], [[intent, *row.values()] for intent, row in sorted(matrix.items())], [16] + [11] * len(DIALECTS))

    by_case = {c.id: c for c in cases}
    detail = lambda r: [r.case_id, by_case[r.case_id].dialect, by_case[r.case_id].intent, by_case[r.case_id].test_type,  # noqa: E731
                        by_case[r.case_id].question, r.answer, ", ".join(r.failed) or "pass", "; ".join(r.missing),
                        "; ".join(r.leaked), r.answer_language, ", ".join(r.behaviors), r.latency_ms, r.error]
    head = ["id", "dialect", "intent", "test type", "question", "answer", "result", "missing content", "leaked",
            "answer language", "behavior", "ms", "error"]
    widths = [10, 11, 13, 12, 44, 54, 18, 22, 16, 10, 12, 8, 20]
    ws = wb.create_sheet("Failures")
    failures = sorted((r for r in run.results if not r.passed), key=lambda r: (r.case_id not in s.critical_failures, r.case_id))
    _table(ws, head, [detail(r) for r in failures], widths)
    for row in ws.iter_rows(min_row=2):
        if row[0].value in s.critical_failures:
            for cell in row:
                cell.fill = BAD
        for cell in row[4:6]:
            cell.alignment = Alignment(wrap_text=True, vertical="top")

    if comparison:
        ws = wb.create_sheet("Comparison")
        c = comparison
        verdict = " (real change)" if c.significant else " (could be chance)"
        _table(ws, ["", "Cases"], [["Fixed", len(c.fixed)], ["Regressed", len(c.regressed)], ["Still failing", len(c.still_failing)],
                                   ["Still passing", c.still_passing], ["McNemar exact p-value", f"{c.p_value:.2g}{verdict}"]],
               [30, 26])
        ws.append([])
        ws.append(["Biggest moves by slice"])
        ws.append(["tag", "value", "cases", "before", "after", "change"])
        for dl in c.deltas:
            ws.append([dl.tag, dl.value, dl.total, round(dl.before, 3), round(dl.after, 3), round(dl.delta, 3)])
        ws.append([])
        ws.append(["Regressed cases"])
        for i in c.regressed:
            ws.append([i, by_case[i].dialect, by_case[i].intent, by_case[i].question])

    ws = wb.create_sheet("All results")
    _table(ws, head, [detail(r) for r in run.results], widths)
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path


def write_markdown(path: str | Path, s: Summary, d: Decision, comparison: Comparison | None = None,
                   cases: list[Case] | None = None, run: Run | None = None, examples: int = 5) -> Path:
    lo, hi = s.ci
    lines = [f"# Chatbot evaluation: {s.label or s.run_id}", "",
             f"**Decision: {d.label}** · bot `{s.bot}` · golden set `{s.dataset}` · {s.total} cases", "",
             f"Pass rate **{s.pass_rate:.1%}** (95% interval {lo:.1%} to {hi:.1%}) · language match {s.language_rate:.1%} · "
             f"critical failures **{len(s.critical_failures)}**", "",
             "## Gates", "", "| Gate | Value | Threshold | Result |", "|---|---|---|---|"]
    lines += [f"| {g.name} | {g.value} | {g.threshold} | {'pass' if g.ok else '**FAIL**'} |" for g in d.gates]
    if d.too_small:
        lines += ["", "Not judged, too few cases: " + "; ".join(d.too_small)]
    lines += ["", "## Why answers failed", "", "| Reason | Cases |", "|---|---|"]
    lines += [f"| {k} | {v} |" for k, v in s.reasons.items() if v]
    for tag in ("dialect", "intent"):
        lines += ["", f"## By {tag}", "", f"| {tag} | Cases | Pass rate | 95% interval |", "|---|---|---|---|"]
        lines += [f"| {sl.value} | {sl.total} | {sl.rate:.1%} | {sl.ci[0]:.1%} to {sl.ci[1]:.1%} |" for sl in s.slices[tag]]
    if comparison:
        c = comparison
        lines += ["", f"## Compared with {c.before}", "",
                  f"Fixed **{len(c.fixed)}**, regressed **{len(c.regressed)}**, still failing {len(c.still_failing)}. "
                  f"McNemar exact p = {c.p_value:.2g} ({'a real change' if c.significant else 'could be chance'}).", "",
                  "| Slice | Before | After | Change |", "|---|---|---|---|"]
        lines += [f"| {dl.tag} = {dl.value} | {dl.before:.1%} | {dl.after:.1%} | {dl.delta:+.1%} |"
                  for dl in c.deltas if abs(dl.delta) >= 0.05]
    if cases and run:
        by_case = {c.id: c for c in cases}
        worst = sorted((r for r in run.results if not r.passed), key=lambda r: (r.case_id not in s.critical_failures, r.case_id))
        lines += ["", "## Examples of failures", "", "| Case | Question | Answer | Why |", "|---|---|---|---|"]
        for r in worst[:examples]:
            q = by_case[r.case_id].question.replace("|", "/")
            a = r.answer.replace("|", "/").replace("\n", " ")[:120]
            lines.append(f"| {r.case_id} ({by_case[r.case_id].dialect}) | {q} | {a} | {', '.join(r.failed)} |")
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text("\n".join(lines) + "\n", encoding="utf-8")
    return path
