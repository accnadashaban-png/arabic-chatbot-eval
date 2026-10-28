"""Human review: export a sheet for reviewers, read their verdicts back, and measure agreement.

If automatic checks (or an LLM judge) disagree often with people, fix the checks before trusting the score.
"""
from __future__ import annotations

import random
from dataclasses import dataclass
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.worksheet.datavalidation import DataValidation

from chatbot_eval.dataset import Case
from chatbot_eval.runner import Run
from chatbot_eval.stats import cohen_kappa

COLUMNS = ["case_id", "dialect", "intent", "question", "answer", "auto_verdict", "auto_reasons", "human_verdict", "comment"]


def sample_for_review(run: Run, size: int = 60, seed: int = 7) -> list[str]:
    """All failures up to half the sample, the rest random passes: reviewers also need to see what 'passed'."""
    rng = random.Random(seed)
    failed = [r.case_id for r in run.results if not r.passed]
    passed = [r.case_id for r in run.results if r.passed]
    take_failed = rng.sample(failed, min(len(failed), size // 2))
    take_passed = rng.sample(passed, min(len(passed), size - len(take_failed)))
    return sorted(take_failed + take_passed)


def export_review(run: Run, cases: list[Case], path: str | Path, ids: list[str] | None = None) -> Path:
    by_case, by_result = {c.id: c for c in cases}, run.by_id()
    ws_book = Workbook()
    ws = ws_book.active
    assert ws is not None
    ws.title = "Review"
    ws.append(COLUMNS)
    for i in ids or sample_for_review(run):
        c, r = by_case[i], by_result[i]
        ws.append([i, c.dialect, c.intent, c.question, r.answer, "pass" if r.passed else "fail", ", ".join(r.failed), "", ""])
    rule = DataValidation(type="list", formula1='"pass,fail"', allow_blank=True)
    ws.add_data_validation(rule)
    rule.add(f"H2:H{ws.max_row}")
    for col, width in zip("ABCDEFGHI", (10, 11, 13, 45, 55, 12, 18, 14, 30), strict=True):
        ws.column_dimensions[col].width = width
    ws.freeze_panes = "A2"
    path = Path(path)
    ws_book.save(path)
    return path


def import_review(path: str | Path) -> dict[str, str]:
    """case_id -> human verdict, for the rows a reviewer filled in."""
    ws = load_workbook(path, read_only=True)["Review"]
    rows = ws.iter_rows(values_only=True)
    header = list(next(rows))
    i, h = header.index("case_id"), header.index("human_verdict")
    verdicts = {}
    for row in rows:
        value = str(row[h] or "").strip().lower()
        if value in ("pass", "fail"):
            verdicts[str(row[i])] = value
    return verdicts


@dataclass
class Agreement:
    reviewed: int
    accuracy: float
    kappa: float
    auto_too_strict: list[str]   # auto said fail, people said pass
    auto_too_lenient: list[str]  # auto said pass, people said fail: the dangerous kind


def agreement(auto: dict[str, str], human: dict[str, str]) -> Agreement:
    ids = sorted(set(auto) & set(human))
    if not ids:
        raise ValueError("no reviewed cases in common")
    a, h = [auto[i] for i in ids], [human[i] for i in ids]
    return Agreement(len(ids), sum(x == y for x, y in zip(a, h, strict=True)) / len(ids), cohen_kappa(a, h),
                     [i for i in ids if auto[i] == "fail" and human[i] == "pass"],
                     [i for i in ids if auto[i] == "pass" and human[i] == "fail"])
